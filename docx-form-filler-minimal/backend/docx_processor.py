"""
Core DOCX template processing logic.

Supports:
- Scalar placeholders like {{item1}}
- Table placeholders like {{table_1}} populated from Excel via pandas
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Mapping

import pandas as pd
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph


PLACEHOLDER_RE = re.compile(r"\{\{([A-Za-z0-9_.-]+)\}\}")
TABLE_PLACEHOLDER_PREFIX = "table_"


@dataclass(frozen=True)
class PlaceholderSummary:
    all_placeholders: list[str]
    scalar_placeholders: list[str]
    table_placeholders: list[str]

    def to_dict(self) -> dict[str, list[str]]:
        return {
            "all_placeholders": self.all_placeholders,
            "scalar_placeholders": self.scalar_placeholders,
            "table_placeholders": self.table_placeholders,
        }


def analyze_docx_placeholders(docx_bytes: bytes) -> PlaceholderSummary:
    """Extract all placeholders from a DOCX file."""
    document = Document(io.BytesIO(docx_bytes))
    all_tokens = _extract_placeholders_from_document(document)

    table_tokens = sorted(
        token for token in all_tokens if token.lower().startswith(TABLE_PLACEHOLDER_PREFIX)
    )
    scalar_tokens = sorted(token for token in all_tokens if token not in table_tokens)

    return PlaceholderSummary(
        all_placeholders=sorted(all_tokens),
        scalar_placeholders=scalar_tokens,
        table_placeholders=table_tokens,
    )


def fill_docx_template(
    docx_bytes: bytes,
    values: Mapping[str, Any] | None = None,
    excel_bytes: bytes | None = None,
) -> bytes:
    """
    Fill a DOCX template with scalar placeholders and optional table placeholders.

    Args:
        docx_bytes: Original DOCX bytes
        values: Placeholder value mapping for scalar fields.
                Keys can be either "item1" or "{{item1}}"
        excel_bytes: Optional Excel bytes for table placeholders ({{table_1}}, etc.)

    Returns:
        Filled DOCX bytes
    """
    document = Document(io.BytesIO(docx_bytes))
    values = values or {}

    scalar_values: dict[str, str] = {}
    table_values_from_json: dict[str, pd.DataFrame] = {}

    for raw_key, raw_value in values.items():
        key = _normalize_placeholder_key(str(raw_key))
        if not key:
            continue

        if key.lower().startswith(TABLE_PLACEHOLDER_PREFIX):
            if isinstance(raw_value, list):
                table_values_from_json[key] = pd.DataFrame(raw_value)
            elif isinstance(raw_value, dict):
                table_values_from_json[key] = pd.DataFrame([raw_value])
            continue

        scalar_values[key] = "" if raw_value is None else str(raw_value)

    # Replace scalar placeholders everywhere in the document.
    for paragraph in _iter_all_paragraphs(document):
        for key, value in scalar_values.items():
            _replace_token_in_paragraph(paragraph, key, value)

    # Determine table placeholders and validate placement support.
    all_table_tokens = {
        token
        for token in _extract_placeholders_from_document(document)
        if token.lower().startswith(TABLE_PLACEHOLDER_PREFIX)
    }
    body_table_tokens = _extract_table_placeholders_from_body(document)

    unsupported_locations = sorted(all_table_tokens - body_table_tokens)
    if unsupported_locations:
        raise ValueError(
            "Table placeholders must be in top-level document body paragraphs: "
            + ", ".join(unsupported_locations)
        )

    if not body_table_tokens:
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    table_map = dict(table_values_from_json)
    if excel_bytes:
        table_map.update(_load_excel_tables(excel_bytes))

    missing = sorted(token for token in body_table_tokens if token not in table_map)
    if missing:
        raise ValueError(
            "Missing table data for placeholders: "
            + ", ".join(missing)
            + ". Provide an Excel file with matching sheets or table_N indices."
        )

    # Replace table placeholders in body paragraphs and insert DataFrame-backed tables.
    for paragraph in list(document.paragraphs):
        table_tokens = _extract_table_tokens_from_text(_paragraph_text(paragraph))
        if not table_tokens:
            continue

        # Reverse insertion to preserve visual order after XML insertions.
        for token in reversed(table_tokens):
            dataframe = table_map[token]
            _replace_token_in_paragraph(paragraph, token, "")
            _insert_dataframe_table_after_paragraph(document, paragraph, dataframe)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _extract_placeholders_from_document(document: DocxDocument) -> set[str]:
    tokens: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        tokens.update(_extract_tokens_from_text(_paragraph_text(paragraph)))
    return tokens


def _extract_table_placeholders_from_body(document: DocxDocument) -> set[str]:
    tokens: set[str] = set()
    for paragraph in document.paragraphs:
        for token in _extract_table_tokens_from_text(_paragraph_text(paragraph)):
            tokens.add(token)
    return tokens


def _extract_tokens_from_text(text: str) -> list[str]:
    return [match.group(1) for match in PLACEHOLDER_RE.finditer(text)]


def _extract_table_tokens_from_text(text: str) -> list[str]:
    ordered: list[str] = []
    for token in _extract_tokens_from_text(text):
        if token.lower().startswith(TABLE_PLACEHOLDER_PREFIX) and token not in ordered:
            ordered.append(token)
    return ordered


def _normalize_placeholder_key(key: str) -> str:
    key = key.strip()
    if key.startswith("{{") and key.endswith("}}"):
        key = key[2:-2].strip()
    return key


def _paragraph_text(paragraph: Paragraph) -> str:
    # Building text from runs avoids paragraph.text side effects.
    run_text = "".join(run.text for run in paragraph.runs)
    return run_text if run_text else paragraph.text


def _replace_token_in_paragraph(paragraph: Paragraph, token: str, replacement: str) -> int:
    """
    Replace placeholder token in a paragraph while preserving run formatting as much as possible.

    Handles cases where a placeholder spans multiple runs.
    """
    placeholder = f"{{{{{token}}}}}"
    replacement = "" if replacement is None else str(replacement)
    replaced_count = 0

    while True:
        if not paragraph.runs:
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, replacement)
                replaced_count += 1
            break

        full_text = "".join(run.text for run in paragraph.runs)
        start = full_text.find(placeholder)
        if start < 0:
            break
        end = start + len(placeholder)

        run_ranges: list[tuple[int, int, int]] = []
        cursor = 0
        for index, run in enumerate(paragraph.runs):
            run_text = run.text or ""
            next_cursor = cursor + len(run_text)
            run_ranges.append((index, cursor, next_cursor))
            cursor = next_cursor

        start_run_index: int | None = None
        end_run_index: int | None = None
        start_run_start = 0
        end_run_start = 0

        for index, run_start, run_end in run_ranges:
            if start_run_index is None and run_start <= start < run_end:
                start_run_index = index
                start_run_start = run_start

            if end_run_index is None and run_start < end <= run_end:
                end_run_index = index
                end_run_start = run_start

        if start_run_index is None or end_run_index is None:
            paragraph.runs[0].text = full_text.replace(placeholder, replacement, 1)
            for run in paragraph.runs[1:]:
                run.text = ""
            replaced_count += 1
            continue

        if start_run_index == end_run_index:
            run = paragraph.runs[start_run_index]
            start_offset = start - start_run_start
            end_offset = end - start_run_start
            run.text = run.text[:start_offset] + replacement + run.text[end_offset:]
        else:
            first_run = paragraph.runs[start_run_index]
            last_run = paragraph.runs[end_run_index]
            start_offset = start - start_run_start
            end_offset = end - end_run_start

            prefix = first_run.text[:start_offset]
            suffix = last_run.text[end_offset:]
            first_run.text = prefix + replacement + suffix

            for i in range(start_run_index + 1, end_run_index + 1):
                paragraph.runs[i].text = ""

        replaced_count += 1

    return replaced_count


def _load_excel_tables(excel_bytes: bytes) -> dict[str, pd.DataFrame]:
    sheets = pd.read_excel(io.BytesIO(excel_bytes), sheet_name=None)
    if not sheets:
        return {}

    table_map: dict[str, pd.DataFrame] = {}
    ordered_sheets = list(sheets.items())

    for idx, (sheet_name, dataframe) in enumerate(ordered_sheets, start=1):
        frame = dataframe.copy()
        frame.columns = [str(col) for col in frame.columns]
        table_map[f"table_{idx}"] = frame
        table_map[sheet_name.strip()] = frame

    # Explicit table_X sheet names should override index mapping.
    for sheet_name, dataframe in ordered_sheets:
        clean_name = sheet_name.strip()
        if clean_name.lower().startswith(TABLE_PLACEHOLDER_PREFIX):
            frame = dataframe.copy()
            frame.columns = [str(col) for col in frame.columns]
            table_map[clean_name] = frame

    return table_map


def _insert_dataframe_table_after_paragraph(
    document: DocxDocument,
    paragraph: Paragraph,
    dataframe: pd.DataFrame,
) -> None:
    frame = dataframe.copy()

    if len(frame.columns) == 0:
        frame = pd.DataFrame({"Value": []})
    else:
        frame.columns = [str(col) for col in frame.columns]

    frame = frame.fillna("")
    columns = list(frame.columns)

    table = document.add_table(rows=1, cols=len(columns))
    try:
        table.style = "Table Grid"
    except Exception:
        # If style doesn't exist in the template, fallback silently.
        pass

    header_cells = table.rows[0].cells
    for idx, column in enumerate(columns):
        header_cells[idx].text = str(column)

    for _, row in frame.iterrows():
        row_cells = table.add_row().cells
        for idx, column in enumerate(columns):
            value = row[column]
            row_cells[idx].text = "" if pd.isna(value) else str(value)

    paragraph._p.addnext(table._tbl)


def _iter_all_paragraphs(document: DocxDocument):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        yield from _iter_table_paragraphs(table)

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for table in section.header.tables:
            yield from _iter_table_paragraphs(table)

        for paragraph in section.footer.paragraphs:
            yield paragraph
        for table in section.footer.tables:
            yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)
